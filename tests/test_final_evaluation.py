import json
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from src.final_evaluation import (
    FinalEvaluationError,
    build_final_rows,
    final_assessment_note,
    final_label,
    load_person_manual_fields,
    main_issue,
    person_final_results,
    representative_rows,
    validation_stats,
)
from src.run_final_evaluation import run_final_evaluation
from src.schemas import FinalPairwiseEvaluation


def automated_row(row_id, person_id="c_human", core=3, f1=0.15, contradicted=0, missing=1,
                   coverage=0.5, strict=0.25, unsupported_rate=0.5, nli_contra=False):
    return {
        "row_id": row_id, "person_id": person_id, "question": f"Q{row_id}?",
        "human_answer": f"H{row_id}", "ai_answer": f"A{row_id}",
        "automated_fidelity_score": core, "automated_fidelity_scale": "1-5",
        "core_fidelity_score": core, "behavior_score": None, "preference_score": None,
        "motivation_score": None, "nuance_score": None,
        "human_claim_count": 4, "aligned_count": 1, "partial_count": 1, "contradicted_count": contradicted,
        "missing_count": missing, "ai_claim_count": 4, "unsupported_ai_count": 2,
        "claim_coverage_rate": coverage, "strict_alignment_rate": strict, "unsupported_rate": unsupported_rate,
        "nli_contradiction_present": nli_contra, "nli_contradiction_count": 1 if nli_contra else 0,
        "strongest_contradiction_probability": 0.9 if nli_contra else None,
        "bertscore_precision": 0.1, "bertscore_recall": 0.2, "bertscore_f1": f1,
    }


def human_row(row_id, person_id="c_human", score=3, contradiction="NO", omission="NO", unsupported="YES",
              note="Some review note."):
    return {
        "review_id": f"r{row_id}", "row_id": row_id, "person_id": person_id,
        "question": f"Q{row_id}?", "human_answer": f"H{row_id}", "ai_answer": f"A{row_id}",
        "human_fidelity_score": score, "human_contradiction": contradiction,
        "human_omission": omission, "human_unsupported_detail": unsupported, "reviewer_note": note,
        "geval_core_fidelity": score, "claim_coverage_rate": 0.5, "strict_alignment_rate": 0.25,
        "step4_contradiction_present": False, "step4_missing_count": 1, "step4_unsupported_rate": 0.5,
        "nli_contradiction_present": False, "nli_contradiction_count": 0,
        "bertscore_precision": 0.1, "bertscore_recall": 0.2, "bertscore_f1": 0.15,
    }


HUMAN_SUMMARY_FIXTURE = {
    "spearman_correlations": {
        "human_fidelity_vs_geval_core_fidelity": {"rho": 0.82, "p_value": 1e-8, "n": 30},
    },
    "contradiction_agreement": {
        "step4_vs_human": {"true_positives": 8, "true_negatives": 19, "false_positives": 0, "false_negatives": 3},
        "nli_vs_human": {"true_positives": 9, "true_negatives": 18, "false_positives": 1, "false_negatives": 2},
    },
    "contradiction_disagreement_rows": {
        "step4_vs_human": [
            {"row_id": 10, "human_contradiction": "YES", "reviewer_note": "n"},
            {"row_id": 11, "human_contradiction": "YES", "reviewer_note": "n"},
        ],
        "nli_vs_human": [
            {"row_id": 7, "human_contradiction": "NO", "reviewer_note": "n"},
            {"row_id": 12, "human_contradiction": "YES", "reviewer_note": "n"},
        ],
    },
}

AUTOMATED_SUMMARY_FIXTURE = {
    "claim_analysis": {"missing_percentage": 50.4, "unsupported_ai_percentage": 78.0},
}


def write_manual_workbook(path: Path, overrides: dict | None = None):
    manual = {
        "c_human": {"consistent_traits": "keeps price", "inconsistent_traits": "reverses behavior",
                    "invented_profile_traits": "adds sustainability", "overall_consistency": "Low",
                    "person_level_note": "Mixed results"},
        "s_human": {"consistent_traits": "keeps brands", "inconsistent_traits": "swaps celebrity",
                    "invented_profile_traits": "adds skin type", "overall_consistency": "Medium",
                    "person_level_note": "Reasonably consistent"},
        "g_human": {"consistent_traits": "keeps supermarket", "inconsistent_traits": "reverses bulk buying",
                    "invented_profile_traits": "adds loyalty", "overall_consistency": "Low",
                    "person_level_note": "Several contradictions"},
    }
    if overrides:
        manual.update(overrides)
    workbook = Workbook()
    workbook.remove(workbook.active)
    fields = ["consistent_traits", "inconsistent_traits", "invented_profile_traits",
              "overall_consistency", "person_level_note"]
    for person_id, values in manual.items():
        sheet = workbook.create_sheet(person_id)
        for row_index, field in enumerate(fields, start=1):
            sheet.cell(row=row_index, column=1, value=field)
            sheet.cell(row=row_index, column=2, value=values[field])
    workbook.save(path)


class FinalLabelTests(unittest.TestCase):
    def test_maps_each_score_to_the_documented_label(self):
        self.assertEqual(final_label(5), "High fidelity")
        self.assertEqual(final_label(4), "Mostly faithful")
        self.assertEqual(final_label(3), "Mixed fidelity")
        self.assertEqual(final_label(2), "Low fidelity")
        self.assertEqual(final_label(1), "Very low fidelity")


class MainIssueTests(unittest.TestCase):
    def test_covers_all_eight_flag_combinations(self):
        self.assertEqual(main_issue(False, False, False), "none / minor")
        self.assertEqual(main_issue(False, False, True), "unsupported elaboration")
        self.assertEqual(main_issue(False, True, False), "important omission")
        self.assertEqual(main_issue(False, True, True), "omission + unsupported elaboration")
        self.assertEqual(main_issue(True, False, False), "direct contradiction")
        self.assertEqual(main_issue(True, False, True), "contradiction + unsupported elaboration")
        self.assertEqual(main_issue(True, True, False), "contradiction + omission")
        self.assertEqual(main_issue(True, True, True), "multiple fidelity issues")


class FinalAssessmentNoteTests(unittest.TestCase):
    def test_appends_overall_label_without_inventing_new_content(self):
        note = final_assessment_note("The AI reverses the behavior", 1)
        self.assertTrue(note.startswith("The AI reverses the behavior."))
        self.assertIn("Overall fidelity: Very low fidelity.", note)


class BuildFinalRowsTests(unittest.TestCase):
    def _write(self, directory: Path, automated, human):
        (directory / "automated_evaluation_by_row.json").write_text(json.dumps({"results": automated}))
        (directory / "human_review_final.json").write_text(
            json.dumps({"results": human, "summary": HUMAN_SUMMARY_FIXTURE})
        )

    def test_score_difference_and_agreement_flags(self):
        with TemporaryDirectory() as directory:
            self._write(
                Path(directory),
                [automated_row(1, core=4)],
                [human_row(1, score=2)],
            )
            with patch("src.final_evaluation.OUTPUT_DIR", Path(directory)):
                [record] = build_final_rows()
        self.assertEqual(record["score_difference"], 2)
        self.assertFalse(record["exact_score_agreement"])
        self.assertFalse(record["within_one_point"])

    def test_exact_agreement_and_within_one_point_true_cases(self):
        with TemporaryDirectory() as directory:
            self._write(
                Path(directory),
                [automated_row(1, core=3), automated_row(2, core=4)],
                [human_row(1, score=3), human_row(2, score=3)],
            )
            with patch("src.final_evaluation.OUTPUT_DIR", Path(directory)):
                records = build_final_rows()
        by_id = {r["row_id"]: r for r in records}
        self.assertTrue(by_id[1]["exact_score_agreement"])
        self.assertTrue(by_id[1]["within_one_point"])
        self.assertFalse(by_id[2]["exact_score_agreement"])
        self.assertTrue(by_id[2]["within_one_point"])

    def test_raises_on_row_id_mismatch_between_sources(self):
        with TemporaryDirectory() as directory:
            self._write(Path(directory), [automated_row(1)], [human_row(2)])
            with patch("src.final_evaluation.OUTPUT_DIR", Path(directory)):
                with self.assertRaises(FinalEvaluationError):
                    build_final_rows()

    def test_raises_on_answer_text_mismatch(self):
        with TemporaryDirectory() as directory:
            automated = [automated_row(1)]
            human = [human_row(1)]
            human[0]["human_answer"] = "a completely different answer"
            self._write(Path(directory), automated, human)
            with patch("src.final_evaluation.OUTPUT_DIR", Path(directory)):
                with self.assertRaises(FinalEvaluationError):
                    build_final_rows()

    def test_every_record_validates_against_the_schema(self):
        with TemporaryDirectory() as directory:
            self._write(Path(directory), [automated_row(1)], [human_row(1)])
            with patch("src.final_evaluation.OUTPUT_DIR", Path(directory)):
                [record] = build_final_rows()
        FinalPairwiseEvaluation.model_validate(record)  # should not raise


class SchemaTests(unittest.TestCase):
    BASE = dict(
        row_id=1, person_id="p", question="Q?", human_answer="H", ai_answer="A",
        automated_fidelity_score=3, automated_fidelity_scale="1-5", behavior_score=None,
        preference_score=None, motivation_score=None, nuance_score=None,
        claim_coverage_rate=0.5, strict_alignment_rate=0.25, contradicted_count=0, missing_count=1,
        unsupported_rate=0.5, nli_contradiction_present=False, bertscore_f1=0.15,
        human_fidelity_score=3, human_contradiction="NO", human_omission="NO",
        human_unsupported_detail="YES", reviewer_note="Note.",
        score_difference=0, exact_score_agreement=True, within_one_point=True,
        final_fidelity_label="Mixed fidelity", main_issue="unsupported elaboration",
        final_assessment_note="Note. Overall fidelity: Mixed fidelity.",
    )

    def test_rejects_wrong_score_difference(self):
        with self.assertRaises(ValidationError):
            FinalPairwiseEvaluation(**{**self.BASE, "score_difference": 1})

    def test_rejects_final_label_not_matching_human_score(self):
        with self.assertRaises(ValidationError):
            FinalPairwiseEvaluation(**{**self.BASE, "final_fidelity_label": "High fidelity"})

    def test_accepts_a_consistent_record(self):
        FinalPairwiseEvaluation(**self.BASE)


class LoadPersonManualFieldsTests(unittest.TestCase):
    def test_reads_consistency_labels_exactly_from_the_workbook(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "person_consistency_review.xlsx"
            write_manual_workbook(path)
            manual = load_person_manual_fields(path)
        self.assertEqual(manual["c_human"]["overall_consistency"], "Low")
        self.assertEqual(manual["s_human"]["overall_consistency"], "Medium")
        self.assertEqual(manual["g_human"]["overall_consistency"], "Low")

    def test_raises_on_unexpected_layout(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "bad.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "c_human"
            sheet.cell(row=1, column=1, value="not_the_expected_field")
            workbook.save(path)
            with self.assertRaises(FinalEvaluationError):
                load_person_manual_fields(path)


class ValidationStatsTests(unittest.TestCase):
    def test_reuses_spearman_result_without_recomputing(self):
        records = [
            {"automated_fidelity_score": 3, "human_fidelity_score": 3, "score_difference": 0,
             "exact_score_agreement": True, "within_one_point": True, "row_id": 1},
            {"automated_fidelity_score": 4, "human_fidelity_score": 2, "score_difference": 2,
             "exact_score_agreement": False, "within_one_point": False, "row_id": 2},
        ]
        result = validation_stats(records, HUMAN_SUMMARY_FIXTURE)
        self.assertEqual(result["spearman_human_fidelity_vs_geval_core"], {"rho": 0.82, "p_value": 1e-8, "n": 30})
        self.assertEqual(result["exact_score_agreement_count"], 1)
        self.assertEqual(result["within_one_point_count"], 1)
        self.assertEqual(result["rows_with_score_difference_over_one"], [2])
        self.assertAlmostEqual(result["mean_absolute_score_difference"], 1.0)


class RepresentativeRowsTests(unittest.TestCase):
    def test_raises_if_a_theme_no_longer_holds(self):
        records = [
            {"row_id": row_id, "person_id": "c_human", "question": "Q", "human_answer": "H", "ai_answer": "A",
             "automated_fidelity_score": 5, "human_fidelity_score": 5, "final_fidelity_label": "High fidelity",
             "main_issue": "none / minor", "human_contradiction": "NO", "human_omission": "NO",
             "nli_contradiction_present": False, "reviewer_note": "note"}
            for row_id in (10, 16, 22, 7, 14, 28)
        ]
        # Row 10's fixture no longer looks like a behavioral reversal (no contradiction, high score).
        with self.assertRaises(FinalEvaluationError):
            representative_rows(records)


class PersonFinalResultsTests(unittest.TestCase):
    def test_uses_manual_labels_exactly_without_recomputing(self):
        records = [
            {"person_id": "c_human", "automated_fidelity_score": 3, "human_fidelity_score": 2,
             "claim_coverage_rate": 0.5, "unsupported_rate": 0.5, "human_contradiction": "YES",
             "human_omission": "NO", "human_unsupported_detail": "YES"},
        ]
        manual = {"c_human": {
            "overall_consistency": "Low", "consistent_traits": "a", "inconsistent_traits": "b",
            "invented_profile_traits": "c", "person_level_note": "note",
        }}
        [summary] = person_final_results(records, manual)
        self.assertEqual(summary["consistency_label"], "Low")
        self.assertNotIn("computed_consistency_score", summary)


class RunFinalEvaluationTests(unittest.TestCase):
    def _write_fixtures(self, directory: Path):
        persons = ["c_human", "s_human", "g_human"]
        # Rows 10/16/22/7/14/28 must satisfy the representative-row theme checks.
        contradiction_overrides = {10: "YES", 16: "YES", 22: "NO", 7: "NO", 14: "YES", 28: "NO"}
        omission_overrides = {22: "NO", 28: "NO"}
        nli_overrides = {7: True}
        core_overrides = {10: 1}

        automated = [
            automated_row(
                i, person_id=persons[(i - 1) % 3], core=core_overrides.get(i, ((i - 1) % 5) + 1),
                nli_contra=nli_overrides.get(i, False),
            )
            for i in range(1, 31)
        ]
        human = [
            human_row(
                i, person_id=persons[(i - 1) % 3], score=((i - 1) % 5) + 1,
                contradiction=contradiction_overrides.get(i, "NO"),
                omission=omission_overrides.get(i, "YES"),
            )
            for i in range(1, 31)
        ]
        (directory / "automated_evaluation_by_row.json").write_text(json.dumps({"results": automated}))
        (directory / "human_review_final.json").write_text(
            json.dumps({"results": human, "summary": HUMAN_SUMMARY_FIXTURE})
        )
        (directory / "automated_evaluation_summary.json").write_text(json.dumps(AUTOMATED_SUMMARY_FIXTURE))
        write_manual_workbook(directory / "person_consistency_review.xlsx")

    def test_full_run_produces_thirty_unique_rows(self):
        with (
            TemporaryDirectory() as directory,
            patch("src.final_evaluation.OUTPUT_DIR", Path(directory)),
            patch("src.run_final_evaluation.OUTPUT_DIR", Path(directory)),
            patch("builtins.print"),
        ):
            output_dir = Path(directory)
            self._write_fixtures(output_dir)
            self.assertEqual(run_final_evaluation(), 0)
            data = json.loads((output_dir / "final_pairwise_evaluation.json").read_text())
        results = data["results"]
        self.assertEqual(len(results), 30)
        self.assertEqual(sorted(r["row_id"] for r in results), list(range(1, 31)))
        self.assertEqual(len({r["row_id"] for r in results}), 30)

    def test_every_row_has_main_issue_and_final_assessment_note(self):
        with (
            TemporaryDirectory() as directory,
            patch("src.final_evaluation.OUTPUT_DIR", Path(directory)),
            patch("src.run_final_evaluation.OUTPUT_DIR", Path(directory)),
            patch("builtins.print"),
        ):
            output_dir = Path(directory)
            self._write_fixtures(output_dir)
            run_final_evaluation()
            data = json.loads((output_dir / "final_pairwise_evaluation.json").read_text())
        for record in data["results"]:
            self.assertTrue(record["main_issue"])
            self.assertTrue(record["final_assessment_note"])

    def test_excel_workbook_has_exactly_three_expected_sheets(self):
        with (
            TemporaryDirectory() as directory,
            patch("src.final_evaluation.OUTPUT_DIR", Path(directory)),
            patch("src.run_final_evaluation.OUTPUT_DIR", Path(directory)),
            patch("builtins.print"),
        ):
            output_dir = Path(directory)
            self._write_fixtures(output_dir)
            run_final_evaluation()
            workbook = load_workbook(output_dir / "final_evaluation.xlsx")
        self.assertEqual(workbook.sheetnames, ["Pairwise Evaluation", "Person Summary", "Overall Summary"])

    def test_person_summary_contains_exactly_three_people(self):
        with (
            TemporaryDirectory() as directory,
            patch("src.final_evaluation.OUTPUT_DIR", Path(directory)),
            patch("src.run_final_evaluation.OUTPUT_DIR", Path(directory)),
            patch("builtins.print"),
        ):
            output_dir = Path(directory)
            self._write_fixtures(output_dir)
            run_final_evaluation()
            summary = json.loads((output_dir / "final_evaluation_summary.json").read_text())
        self.assertEqual(len(summary["person_summary"]), 3)
        self.assertEqual({p["person_id"] for p in summary["person_summary"]}, {"c_human", "s_human", "g_human"})

    def test_does_not_modify_previous_step_outputs(self):
        with (
            TemporaryDirectory() as directory,
            patch("src.final_evaluation.OUTPUT_DIR", Path(directory)),
            patch("src.run_final_evaluation.OUTPUT_DIR", Path(directory)),
            patch("builtins.print"),
        ):
            output_dir = Path(directory)
            self._write_fixtures(output_dir)
            before = {
                name: (output_dir / name).read_bytes()
                for name in (
                    "automated_evaluation_by_row.json", "human_review_final.json",
                    "automated_evaluation_summary.json", "person_consistency_review.xlsx",
                )
            }
            run_final_evaluation()
            for name, content in before.items():
                self.assertEqual((output_dir / name).read_bytes(), content)

    def test_output_filenames_are_exact(self):
        with (
            TemporaryDirectory() as directory,
            patch("src.final_evaluation.OUTPUT_DIR", Path(directory)),
            patch("src.run_final_evaluation.OUTPUT_DIR", Path(directory)),
            patch("builtins.print"),
        ):
            output_dir = Path(directory)
            self._write_fixtures(output_dir)
            run_final_evaluation()
            for name in (
                "final_pairwise_evaluation.json", "final_pairwise_evaluation.csv",
                "final_evaluation_summary.json", "final_evaluation_summary.md", "final_evaluation.xlsx",
            ):
                self.assertTrue((output_dir / name).exists(), name)


if __name__ == "__main__":
    unittest.main()
