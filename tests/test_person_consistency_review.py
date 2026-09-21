"""Tests for the person-level consistency review workbook."""
import json
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import load_workbook

from src.person_consistency_review import (
    EXPECTED_PERSON_IDS,
    MANUAL_REVIEW_FIELDS,
    ROW_COLUMNS,
    PersonConsistencyError,
    build_workbook,
    group_by_person,
    person_summary,
    validate_grouping,
)
from src.run_person_consistency_review import run_person_consistency_review


def make_record(row_id: int, person_id: str, **overrides) -> dict:
    base = {
        "review_id": f"r{row_id}", "row_id": row_id, "person_id": person_id,
        "question": f"Q{row_id}?", "human_answer": f"H{row_id}", "ai_answer": f"A{row_id}",
        "human_fidelity_score": 3, "human_contradiction": "NO", "human_omission": "NO",
        "human_unsupported_detail": "YES", "reviewer_note": f"Note {row_id}.",
    }
    return {**base, **overrides}


def full_thirty_records() -> list[dict]:
    persons = ["c_human", "s_human", "g_human"]
    return [make_record(i, persons[(i - 1) % 3]) for i in range(1, 31)]


class GroupByPersonTests(unittest.TestCase):
    def test_groups_and_sorts_by_row_id_within_each_person(self):
        records = [
            make_record(3, "c_human"), make_record(1, "c_human"), make_record(2, "s_human"),
        ]
        grouped = group_by_person(records)
        self.assertEqual([r["row_id"] for r in grouped["c_human"]], [1, 3])
        self.assertEqual([r["row_id"] for r in grouped["s_human"]], [2])


class ValidateGroupingTests(unittest.TestCase):
    def test_accepts_exactly_thirty_rows_across_the_three_expected_people(self):
        records = full_thirty_records()
        grouped = group_by_person(records)
        validate_grouping(records, grouped)  # should not raise

    def test_rejects_wrong_total_row_count(self):
        records = full_thirty_records()[:29]
        grouped = group_by_person(records)
        with self.assertRaisesRegex(PersonConsistencyError, "exactly 30"):
            validate_grouping(records, grouped)

    def test_rejects_unexpected_person_ids(self):
        records = [make_record(i, "unknown_person") for i in range(1, 31)]
        grouped = group_by_person(records)
        with self.assertRaisesRegex(PersonConsistencyError, "person_ids"):
            validate_grouping(records, grouped)

    def test_rejects_missing_expected_person(self):
        records = [make_record(i, "c_human" if i <= 15 else "s_human") for i in range(1, 31)]
        grouped = group_by_person(records)
        with self.assertRaisesRegex(PersonConsistencyError, "person_ids"):
            validate_grouping(records, grouped)


class BuildWorkbookTests(unittest.TestCase):
    def test_creates_exactly_three_sheets_named_by_person_id(self):
        grouped = group_by_person(full_thirty_records())
        workbook = build_workbook(grouped)
        self.assertEqual(set(workbook.sheetnames), EXPECTED_PERSON_IDS)

    def test_manual_review_fields_are_blank_and_at_the_top(self):
        grouped = group_by_person(full_thirty_records())
        workbook = build_workbook(grouped)
        sheet = workbook["c_human"]
        for row_index, field in enumerate(MANUAL_REVIEW_FIELDS, start=1):
            self.assertEqual(sheet.cell(row=row_index, column=1).value, field)
            self.assertIsNone(sheet.cell(row=row_index, column=2).value)

    def test_row_level_table_preserves_answer_text_and_row_ids(self):
        grouped = group_by_person([
            make_record(1, "c_human", human_answer="Exact human text.", ai_answer="Exact AI text."),
        ])
        workbook = build_workbook(grouped)
        sheet = workbook["c_human"]
        header_row = len(MANUAL_REVIEW_FIELDS) + 2
        headers = [sheet.cell(row=header_row, column=i).value for i in range(1, len(ROW_COLUMNS) + 1)]
        self.assertEqual(headers, ROW_COLUMNS)
        data_row = [sheet.cell(row=header_row + 1, column=i).value for i in range(1, len(ROW_COLUMNS) + 1)]
        self.assertEqual(data_row[0], 1)
        self.assertEqual(data_row[ROW_COLUMNS.index("human_answer")], "Exact human text.")
        self.assertEqual(data_row[ROW_COLUMNS.index("ai_answer")], "Exact AI text.")

    def test_each_row_appears_exactly_once_attached_to_its_own_person(self):
        grouped = group_by_person(full_thirty_records())
        workbook = build_workbook(grouped)
        header_row = len(MANUAL_REVIEW_FIELDS) + 2
        seen_row_ids = []
        for person_id in workbook.sheetnames:
            sheet = workbook[person_id]
            row = header_row + 1
            while sheet.cell(row=row, column=1).value is not None:
                seen_row_ids.append(sheet.cell(row=row, column=1).value)
                row += 1
        self.assertEqual(sorted(seen_row_ids), list(range(1, 31)))
        self.assertEqual(len(seen_row_ids), len(set(seen_row_ids)))


class PersonSummaryTests(unittest.TestCase):
    def test_computes_counts_and_mean_without_a_consistency_rating(self):
        records = [
            make_record(1, "c_human", human_fidelity_score=2, human_contradiction="YES", human_omission="NO"),
            make_record(2, "c_human", human_fidelity_score=4, human_contradiction="NO", human_omission="YES"),
        ]
        grouped = group_by_person(records)
        [summary] = person_summary(grouped)
        self.assertEqual(summary["person_id"], "c_human")
        self.assertEqual(summary["row_ids"], [1, 2])
        self.assertEqual(summary["number_of_rows"], 2)
        self.assertAlmostEqual(summary["mean_human_fidelity_score"], 3.0)
        self.assertEqual(summary["human_contradiction_count"], 1)
        self.assertEqual(summary["human_omission_count"], 1)
        self.assertEqual(summary["human_unsupported_detail_count"], 2)
        self.assertNotIn("overall_consistency", summary)


class RunPersonConsistencyReviewTests(unittest.TestCase):
    def test_writes_expected_files_with_blank_manual_fields(self):
        records = full_thirty_records()
        with (
            TemporaryDirectory() as directory,
            patch("src.run_person_consistency_review.OUTPUT_DIR", Path(directory)),
            patch(
                "src.run_person_consistency_review.load_human_review_records",
                return_value=records,
            ),
            patch("builtins.print"),
        ):
            self.assertEqual(run_person_consistency_review(), 0)
            output_dir = Path(directory)
            xlsx_path = output_dir / "person_consistency_review.xlsx"
            json_path = output_dir / "person_consistency_context.json"
            self.assertTrue(xlsx_path.exists())
            self.assertTrue(json_path.exists())

            workbook = load_workbook(xlsx_path)
            self.assertEqual(set(workbook.sheetnames), EXPECTED_PERSON_IDS)
            for field_row, field in enumerate(MANUAL_REVIEW_FIELDS, start=1):
                self.assertIsNone(workbook["c_human"].cell(row=field_row, column=2).value)

            summary = json.loads(json_path.read_text())
            self.assertEqual(len(summary["people"]), 3)
            for person in summary["people"]:
                self.assertNotIn("overall_consistency", person)
                self.assertEqual(person["number_of_rows"], 10)

    def test_does_not_modify_frozen_step3_to_7b_outputs(self):
        records = full_thirty_records()
        with (
            TemporaryDirectory() as directory,
            patch("src.run_person_consistency_review.OUTPUT_DIR", Path(directory)),
            patch(
                "src.run_person_consistency_review.load_human_review_records",
                return_value=records,
            ),
            patch("builtins.print"),
        ):
            output_dir = Path(directory)
            frozen_files = [
                "geval_results_final.json", "claim_alignment_final.json",
                "nli_validation_final.json", "bertscore_final.json", "human_review_final.json",
            ]
            before = {}
            for name in frozen_files:
                path = output_dir / name
                path.write_text(f'{{"marker": "{name}"}}')
                before[name] = path.read_text()
            run_person_consistency_review()
            for name in frozen_files:
                self.assertEqual((output_dir / name).read_text(), before[name])


if __name__ == "__main__":
    unittest.main()
