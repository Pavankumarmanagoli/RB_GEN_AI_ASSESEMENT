"""Group human review rows by person and build the consistency review workbook."""
import json
import statistics
from collections import defaultdict

from openpyxl import Workbook

from src.config import OUTPUT_DIR

EXPECTED_PERSON_IDS = {"c_human", "s_human", "g_human"}

# Person-level fields for the reviewer to fill in by hand; never auto-filled.
MANUAL_REVIEW_FIELDS = [
    "consistent_traits",
    "inconsistent_traits",
    "invented_profile_traits",
    "overall_consistency",
    "person_level_note",
]
ROW_COLUMNS = [
    "row_id", "question", "human_answer", "ai_answer",
    "human_fidelity_score", "human_contradiction", "human_omission", "human_unsupported_detail",
    "reviewer_note",
]


class PersonConsistencyError(ValueError):
    """Raised for invalid person-review groupings."""


def load_human_review_records() -> list[dict]:
    data = json.loads((OUTPUT_DIR / "human_review_final.json").read_text(encoding="utf-8"))
    return data["results"]


def group_by_person(records: list[dict]) -> dict[str, list[dict]]:
    grouped = defaultdict(list)
    for record in records:
        grouped[record["person_id"]].append(record)
    return {person_id: sorted(rows, key=lambda r: r["row_id"]) for person_id, rows in grouped.items()}


def validate_grouping(records: list[dict], grouped: dict[str, list[dict]]) -> None:
    if len(records) != 30:
        raise PersonConsistencyError(f"Expected exactly 30 human review rows, found {len(records)}.")
    if set(grouped) != EXPECTED_PERSON_IDS:
        raise PersonConsistencyError(f"Expected exactly person_ids {sorted(EXPECTED_PERSON_IDS)}, found {sorted(grouped)}.")

    all_row_ids = [record["row_id"] for rows in grouped.values() for record in rows]
    if len(all_row_ids) != 30:
        raise PersonConsistencyError(f"Expected exactly 30 grouped rows, found {len(all_row_ids)}.")
    duplicates = [row_id for row_id in set(all_row_ids) if all_row_ids.count(row_id) > 1]
    if duplicates:
        raise PersonConsistencyError(f"Duplicate row_id found across person groups: {sorted(duplicates)}.")


def build_workbook(grouped: dict[str, list[dict]]) -> Workbook:
    """Build one sheet per person with a blank manual-review block above the row-level evidence."""
    workbook = Workbook()
    workbook.remove(workbook.active)

    for person_id in sorted(grouped):
        sheet = workbook.create_sheet(title=person_id)
        for row_index, field in enumerate(MANUAL_REVIEW_FIELDS, start=1):
            sheet.cell(row=row_index, column=1, value=field)
            sheet.cell(row=row_index, column=2, value=None)

        header_row = len(MANUAL_REVIEW_FIELDS) + 2
        for column_index, column_name in enumerate(ROW_COLUMNS, start=1):
            sheet.cell(row=header_row, column=column_index, value=column_name)

        for offset, record in enumerate(grouped[person_id], start=1):
            for column_index, column_name in enumerate(ROW_COLUMNS, start=1):
                sheet.cell(row=header_row + offset, column=column_index, value=record[column_name])

    return workbook


def person_summary(grouped: dict[str, list[dict]]) -> list[dict]:
    summaries = []
    for person_id in sorted(grouped):
        rows = grouped[person_id]
        scores = [row["human_fidelity_score"] for row in rows]
        summaries.append({
            "person_id": person_id,
            "row_ids": [row["row_id"] for row in rows],
            "number_of_rows": len(rows),
            "mean_human_fidelity_score": statistics.mean(scores),
            "human_contradiction_count": sum(1 for row in rows if row["human_contradiction"] == "YES"),
            "human_omission_count": sum(1 for row in rows if row["human_omission"] == "YES"),
            "human_unsupported_detail_count": sum(1 for row in rows if row["human_unsupported_detail"] == "YES"),
        })
    return summaries
