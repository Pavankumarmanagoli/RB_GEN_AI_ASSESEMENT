import json
import re
import statistics

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from src.config import OUTPUT_DIR
from src.schemas import FINAL_FIDELITY_LABELS, FinalPairwiseEvaluation

MANUAL_REVIEW_FIELDS = [
    "consistent_traits", "inconsistent_traits", "invented_profile_traits",
    "overall_consistency", "person_level_note",
]

# Presentation-layer only: maps the technical FINAL_FIDELITY_LABELS values to the
# business-facing "match" wording used in reviewer-facing prose, Excel, and markdown.
# The underlying final_fidelity_label field/value and FINAL_FIDELITY_LABELS mapping in
# src/schemas.py are untouched; this is purely a display-time relabeling.
BUSINESS_MATCH_LABELS = {
    "High fidelity": "Very strong match",
    "Mostly faithful": "Strong match",
    "Mixed fidelity": "Partial match",
    "Low fidelity": "Weak match",
    "Very low fidelity": "Very weak match",
}


def business_match_label(technical_label: str) -> str:
    return BUSINESS_MATCH_LABELS[technical_label]


# Presentation-layer only: softens leftover "fidelity" wording in prose that gets
# displayed in Excel/markdown. The underlying main_issue and final_assessment_note
# field values are untouched everywhere else (e.g. final_pairwise_evaluation.json).
# Sorted longest-key-first so "Very low fidelity" is replaced before "Low fidelity".
_BUSINESS_TEXT_REPLACEMENTS = [
    ("multiple fidelity issues", "multiple issues"),
    ("Overall fidelity:", "Overall match:"),
] + sorted(BUSINESS_MATCH_LABELS.items(), key=lambda item: -len(item[0]))


def business_facing_text(text: str) -> str:
    result = text
    for old, new in _BUSINESS_TEXT_REPLACEMENTS:
        result = result.replace(old, new)
    return result

# Verified against the frozen data: theme, and the condition each row must satisfy for
# that theme to still hold, checked at build time rather than assumed.
REPRESENTATIVE_ROW_THEMES = [
    (10, "Behavioral reversal", lambda r: r["human_contradiction"] == "YES" and r["automated_fidelity_score"] <= 2),
    (16, "Unsupported switching reason", lambda r: r["human_contradiction"] == "YES"),
    (22, "Correct choice, heavy unsupported elaboration",
     lambda r: r["human_contradiction"] == "NO" and r["human_omission"] == "NO"),
    (7, "NLI false positive", lambda r: r["human_contradiction"] == "NO" and r["nli_contradiction_present"]),
    (14, "Celebrity preference mismatch", lambda r: r["human_contradiction"] == "YES"),
    (28, "Broad theme preserved, heavy elaboration",
     lambda r: r["human_contradiction"] == "NO" and r["human_omission"] == "NO"),
]


class FinalEvaluationError(ValueError):
    pass


def load_automated_rows() -> dict:
    data = json.loads((OUTPUT_DIR / "automated_evaluation_by_row.json").read_text(encoding="utf-8"))
    return {record["row_id"]: record for record in data["results"]}


def load_human_review() -> tuple[dict, dict]:
    data = json.loads((OUTPUT_DIR / "human_review_final.json").read_text(encoding="utf-8"))
    return {record["row_id"]: record for record in data["results"]}, data["summary"]


def load_automated_summary() -> dict:
    return json.loads((OUTPUT_DIR / "automated_evaluation_summary.json").read_text(encoding="utf-8"))


def load_person_manual_fields(path=None) -> dict:
    path = path or (OUTPUT_DIR / "person_consistency_review.xlsx")
    workbook = load_workbook(path, data_only=True)
    manual_by_person = {}
    for person_id in workbook.sheetnames:
        sheet = workbook[person_id]
        fields = {}
        for row_index, field_name in enumerate(MANUAL_REVIEW_FIELDS, start=1):
            label = sheet.cell(row=row_index, column=1).value
            if label != field_name:
                raise FinalEvaluationError(
                    f"Unexpected manual review layout in sheet '{person_id}' at row {row_index}: "
                    f"expected {field_name!r}, found {label!r}."
                )
            fields[field_name] = sheet.cell(row=row_index, column=2).value
        manual_by_person[person_id] = fields
    return manual_by_person


def _safe_mean(values: list) -> float | None:
    filtered = [value for value in values if value is not None]
    return statistics.mean(filtered) if filtered else None


def final_label(human_fidelity_score: int) -> str:
    return FINAL_FIDELITY_LABELS[human_fidelity_score]


def main_issue(contradiction: bool, omission: bool, unsupported: bool) -> str:
    """Deterministic mapping from the three Human-review flags to one of the fixed
    issue categories; no heuristics or scoring involved."""
    if contradiction and omission and unsupported:
        return "multiple fidelity issues"
    if contradiction and omission:
        return "contradiction + omission"
    if contradiction and unsupported:
        return "contradiction + unsupported elaboration"
    if omission and unsupported:
        return "omission + unsupported elaboration"
    if contradiction:
        return "direct contradiction"
    if omission:
        return "important omission"
    if unsupported:
        return "unsupported elaboration"
    return "none / minor"


def _fix_reviewer_note_grammar(note: str) -> str:
    """Corrects a stray mid-sentence capitalization in a couple of source review notes
    ('... and It also ...' -> '... and it also ...'). Wording and meaning are unchanged;
    only the capital letter is fixed. The frozen Step 7B reviewer_note itself is never
    modified — this only affects how Step 8B renders it."""
    return re.sub(r"\band It\b", "and it", note)


def final_assessment_note(reviewer_note: str, human_fidelity_score: int) -> str:
    """Built only from the existing, human-authored reviewer_note plus a deterministic
    label lookup — no new LLM judgment or numeric score is generated here."""
    note = _fix_reviewer_note_grammar(reviewer_note.strip())
    if note and not note.endswith((".", "!", "?")):
        note += "."
    return f"{note} Overall fidelity: {final_label(human_fidelity_score)}."


def build_final_rows() -> list[dict]:
    automated_by_row = load_automated_rows()
    human_by_row, _ = load_human_review()

    if set(automated_by_row) != set(human_by_row):
        raise FinalEvaluationError(
            "Automated and Human review row IDs do not match: "
            f"automated only={set(automated_by_row) - set(human_by_row)}, "
            f"human only={set(human_by_row) - set(automated_by_row)}."
        )

    records = []
    for row_id in sorted(automated_by_row):
        automated = automated_by_row[row_id]
        human = human_by_row[row_id]

        if automated["human_answer"] != human["human_answer"] or automated["ai_answer"] != human["ai_answer"]:
            raise FinalEvaluationError(f"Row {row_id}: Human/AI answer text differs between the two source files.")
        if automated["person_id"] != human["person_id"]:
            raise FinalEvaluationError(f"Row {row_id}: person_id differs between the two source files.")

        automated_score = automated["automated_fidelity_score"]
        human_score = human["human_fidelity_score"]
        score_difference = automated_score - human_score

        contradiction = human["human_contradiction"] == "YES"
        omission = human["human_omission"] == "YES"
        unsupported = human["human_unsupported_detail"] == "YES"

        record = FinalPairwiseEvaluation(
            row_id=row_id,
            person_id=automated["person_id"],
            question=automated["question"],
            human_answer=automated["human_answer"],
            ai_answer=automated["ai_answer"],
            automated_fidelity_score=automated_score,
            automated_fidelity_scale=automated["automated_fidelity_scale"],
            behavior_score=automated["behavior_score"],
            preference_score=automated["preference_score"],
            motivation_score=automated["motivation_score"],
            nuance_score=automated["nuance_score"],
            claim_coverage_rate=automated["claim_coverage_rate"],
            strict_alignment_rate=automated["strict_alignment_rate"],
            contradicted_count=automated["contradicted_count"],
            missing_count=automated["missing_count"],
            unsupported_rate=automated["unsupported_rate"],
            nli_contradiction_present=automated["nli_contradiction_present"],
            bertscore_f1=automated["bertscore_f1"],
            human_fidelity_score=human_score,
            human_contradiction=human["human_contradiction"],
            human_omission=human["human_omission"],
            human_unsupported_detail=human["human_unsupported_detail"],
            reviewer_note=human["reviewer_note"],
            score_difference=score_difference,
            exact_score_agreement=(score_difference == 0),
            within_one_point=(abs(score_difference) <= 1),
            final_fidelity_label=final_label(human_score),
            main_issue=main_issue(contradiction, omission, unsupported),
            final_assessment_note=final_assessment_note(human["reviewer_note"], human_score),
        )
        records.append(record.model_dump())

    return records


def validation_stats(records: list[dict], human_summary: dict) -> dict:
    """Reuses the frozen Step 7B Spearman result verbatim rather than recomputing it."""
    automated_scores = [record["automated_fidelity_score"] for record in records]
    human_scores = [record["human_fidelity_score"] for record in records]
    abs_diffs = [abs(record["score_difference"]) for record in records]
    n = len(records)
    exact = sum(1 for record in records if record["exact_score_agreement"])
    within_one = sum(1 for record in records if record["within_one_point"])
    spearman = human_summary["spearman_correlations"]["human_fidelity_vs_geval_core_fidelity"]

    return {
        "automated_fidelity_mean": statistics.mean(automated_scores),
        "automated_fidelity_median": statistics.median(automated_scores),
        "human_fidelity_mean": statistics.mean(human_scores),
        "human_fidelity_median": statistics.median(human_scores),
        "mean_difference": statistics.mean(automated_scores) - statistics.mean(human_scores),
        "mean_absolute_score_difference": statistics.mean(abs_diffs),
        "exact_score_agreement_count": exact,
        "exact_score_agreement_percentage": exact / n * 100,
        "within_one_point_count": within_one,
        "within_one_point_percentage": within_one / n * 100,
        "rows_with_score_difference_over_one": sorted(
            record["row_id"] for record in records if abs(record["score_difference"]) > 1
        ),
        "spearman_human_fidelity_vs_geval_core": {
            "rho": spearman["rho"], "p_value": spearman["p_value"], "n": spearman["n"],
        },
    }


def contradiction_validation(human_summary: dict) -> dict:
    """Reuses the frozen Step 7B confusion matrices and disagreement rows verbatim;
    no new contradiction judgment is made here."""
    agreement = human_summary["contradiction_agreement"]
    disagreement = human_summary["contradiction_disagreement_rows"]
    step4 = agreement["step4_vs_human"]
    nli = agreement["nli_vs_human"]

    return {
        "step4_vs_human": step4,
        "nli_vs_human": nli,
        "contradictions_correctly_detected": {"step4": step4["true_positives"], "nli": nli["true_positives"]},
        "contradictions_missed_by_step4": {
            "count": step4["false_negatives"],
            "rows": sorted(entry["row_id"] for entry in disagreement["step4_vs_human"]),
        },
        "contradictions_missed_by_nli": {
            "count": nli["false_negatives"],
            "rows": sorted(
                entry["row_id"] for entry in disagreement["nli_vs_human"] if entry["human_contradiction"] == "YES"
            ),
        },
        "nli_false_positives": {
            "count": nli["false_positives"],
            "rows": sorted(
                entry["row_id"] for entry in disagreement["nli_vs_human"] if entry["human_contradiction"] == "NO"
            ),
        },
    }


# These additions are unverified against the Human reference, not proven fabrications;
# "invented" overstates that certainty, so person-level text uses "unsupported" instead.
# The invented_profile_traits field name is kept as-is for output-schema compatibility.
_WORDING_SOFTENERS = [
    ("invented profile details", "unsupported profile details"),
    ("Invented profile details", "Unsupported profile details"),
    ("invented profile traits", "unsupported profile traits"),
    ("Invented profile traits", "Unsupported profile traits"),
    ("invented personal details", "unsupported personal details"),
    ("Invented personal details", "Unsupported personal details"),
    ("invented details", "unsupported personal details"),
    ("Invented details", "Unsupported personal details"),
]


def _soften_wording(text: str | None) -> str | None:
    if not text:
        return text
    softened = text
    for old, new in _WORDING_SOFTENERS:
        softened = softened.replace(old, new)
    return softened


def person_final_conclusion(person_id: str, automated_mean: float, human_mean: float, manual: dict) -> str:
    note = (manual["person_level_note"] or "").strip().rstrip(".")
    return (
        f"{person_id}'s simulated answers show an Automated Match Score of {automated_mean:.2f}/5 and a "
        f"Human-Reviewed Match Score of {human_mean:.2f}/5, with overall person-level consistency rated "
        f"{manual['overall_consistency']}. {note}."
    )


def person_final_results(records: list[dict], manual_by_person: dict) -> list[dict]:
    summaries = []
    for person_id in sorted({record["person_id"] for record in records}):
        rows = [record for record in records if record["person_id"] == person_id]
        raw_manual = manual_by_person[person_id]
        manual = {
            "overall_consistency": raw_manual["overall_consistency"],
            "consistent_traits": _soften_wording(raw_manual["consistent_traits"]),
            "inconsistent_traits": _soften_wording(raw_manual["inconsistent_traits"]),
            "invented_profile_traits": _soften_wording(raw_manual["invented_profile_traits"]),
            "person_level_note": _soften_wording(raw_manual["person_level_note"]),
        }
        automated_mean = statistics.mean([row["automated_fidelity_score"] for row in rows])
        human_mean = statistics.mean([row["human_fidelity_score"] for row in rows])
        summaries.append({
            "person_id": person_id,
            "number_of_rows": len(rows),
            "automated_fidelity_mean": automated_mean,
            "human_fidelity_mean": human_mean,
            "mean_claim_coverage_rate": _safe_mean([row["claim_coverage_rate"] for row in rows]),
            "mean_unsupported_rate": _safe_mean([row["unsupported_rate"] for row in rows]),
            "human_contradiction_count": sum(1 for row in rows if row["human_contradiction"] == "YES"),
            "human_omission_count": sum(1 for row in rows if row["human_omission"] == "YES"),
            "human_unsupported_detail_count": sum(1 for row in rows if row["human_unsupported_detail"] == "YES"),
            "consistency_label": manual["overall_consistency"],
            "consistent_traits": manual["consistent_traits"],
            "inconsistent_traits": manual["inconsistent_traits"],
            "invented_profile_traits": manual["invented_profile_traits"],
            "person_level_note": manual["person_level_note"],
            "person_final_conclusion": person_final_conclusion(person_id, automated_mean, human_mean, manual),
        })
    return summaries


def _short_summary(text: str, limit: int = 140) -> str:
    text = text.strip()
    return text if len(text) <= limit else text[:limit].rstrip() + "…"


def representative_rows(records: list[dict]) -> list[dict]:
    """Selects a fixed set of rows chosen for how clearly they illustrate a fidelity
    pattern to a business reviewer; each theme is re-verified against the joined data
    before use, so a row is never described a way the data no longer supports."""
    by_row = {record["row_id"]: record for record in records}
    examples = []
    for row_id, theme, condition in REPRESENTATIVE_ROW_THEMES:
        record = by_row[row_id]
        if not condition(record):
            raise FinalEvaluationError(f"Representative row {row_id} no longer supports '{theme}'.")
        examples.append({
            "row_id": row_id,
            "person_id": record["person_id"],
            "question": record["question"],
            "human_answer_summary": _short_summary(record["human_answer"]),
            "ai_answer_summary": _short_summary(record["ai_answer"]),
            "automated_fidelity_score": record["automated_fidelity_score"],
            "human_fidelity_score": record["human_fidelity_score"],
            "final_fidelity_label": record["final_fidelity_label"],
            "main_issue": record["main_issue"],
            "explanation": f"{theme}: {_fix_reviewer_note_grammar(record['reviewer_note'])}",
        })
    return examples


def final_business_findings(records: list[dict], validation: dict, human_rates: dict, persons: list[dict]) -> list[str]:
    low_persons = [p["person_id"] for p in persons if p["consistency_label"] == "Low"]
    medium_persons = [p["person_id"] for p in persons if p["consistency_label"] == "Medium"]
    consistency_sentence = "Person-level consistency "
    parts = []
    if low_persons:
        parts.append(f"remains weak for {', '.join(low_persons)} (Low)")
    if medium_persons:
        parts.append(f"is only moderate for {', '.join(medium_persons)} (Medium)")
    consistency_sentence += " and ".join(parts) + ", based on the manual cross-row consistency review." if parts \
        else "varies across the three simulated people, based on the manual cross-row consistency review."

    return [
        f"Automated and Human-reviewed match scores broadly agree: the Automated Match Score averages "
        f"{validation['automated_fidelity_mean']:.2f}/5 versus a Human-Reviewed Match Score of "
        f"{validation['human_fidelity_mean']:.2f}/5, matching exactly on "
        f"{validation['exact_score_agreement_percentage']:.0f}% of rows and within one point on "
        f"{validation['within_one_point_percentage']:.0f}%.",
        "Broad themes (what a person buys, cares about, or discovers) are usually preserved, but "
        "person-specific details, quantities, and named preferences are frequently lost or replaced.",
        f"Omissions and unsupported elaboration are common: the human reviewer flagged omission in "
        f"{human_rates['human_omission_rate']:.0f}% of rows and unsupported detail in "
        f"{human_rates['human_unsupported_detail_rate']:.0f}% of rows.",
        f"Human review identified contradictions in {human_rates['human_contradiction_rate']:.1f}% of rows, "
        "including several clear behavioral or preference reversals.",
        consistency_sentence,
    ]


def build_overall_conclusion(validation: dict, human_rates: dict, dataset_size: int) -> str:
    spearman = validation["spearman_human_fidelity_vs_geval_core"]
    if validation["within_one_point_percentage"] == 100.0:
        agreement_clause = f"differ by no more than one point across all {dataset_size} rows"
    else:
        agreement_clause = (
            f"differ by no more than one point on {validation['within_one_point_percentage']:.0f}% of rows"
        )
    return (
        f"Across the {dataset_size} reviewed answers, the AI simulations show only a partial match to the "
        f"real Human responses, with an Automated Match Score of {validation['automated_fidelity_mean']:.2f}/5 "
        f"and a Human-Reviewed Match Score of {validation['human_fidelity_mean']:.2f}/5. The simulations often "
        "preserve parts of the broad theme of an answer, but frequently omit person-specific details and "
        f"introduce unsupported personal elaboration; Human review identified contradictions in "
        f"{human_rates['human_contradiction_rate']:.1f}% of rows, including several clear behavioral or "
        f"preference reversals. Automated and Human judgments show strong rank agreement (Spearman ρ = "
        f"{spearman['rho']:.2f}) and {agreement_clause}, indicating that automated evaluation can serve as "
        "a useful screening signal. Human review remains important for identifying meaningful "
        "reversals, determining which omissions matter, and resolving borderline cases where automated "
        "methods disagree."
    )


def dataset_final_summary(records: list[dict], human_summary: dict, manual_by_person: dict, automated_summary: dict) -> dict:
    validation = validation_stats(records, human_summary)
    contradiction = contradiction_validation(human_summary)
    persons = person_final_results(records, manual_by_person)
    claims = automated_summary["claim_analysis"]

    n = len(records)
    human_rates = {
        "human_contradiction_rate": sum(1 for r in records if r["human_contradiction"] == "YES") / n * 100,
        "human_omission_rate": sum(1 for r in records if r["human_omission"] == "YES") / n * 100,
        "human_unsupported_detail_rate": sum(1 for r in records if r["human_unsupported_detail"] == "YES") / n * 100,
        "missing_human_claim_rate": claims["missing_percentage"],
        "unsupported_ai_claim_rate": claims["unsupported_ai_percentage"],
    }

    return {
        "dataset_size": n,
        "automated_fidelity_mean": validation["automated_fidelity_mean"],
        "human_fidelity_mean": validation["human_fidelity_mean"],
        "validation": validation,
        "contradiction_validation": contradiction,
        "failure_evidence": human_rates,
        "person_consistency": {person["person_id"]: person["consistency_label"] for person in persons},
        "person_summary": persons,
        "representative_examples": representative_rows(records),
        "final_business_findings": final_business_findings(records, validation, human_rates, persons),
        "overall_conclusion": build_overall_conclusion(validation, human_rates, n),
    }


def build_markdown_summary(summary: dict) -> str:
    validation = summary["validation"]
    contradiction = summary["contradiction_validation"]
    failure = summary["failure_evidence"]
    spearman = validation["spearman_human_fidelity_vs_geval_core"]

    lines = [
        "# Final Evaluation Summary",
        "",
        "## Automated vs Human-Reviewed Match",
        "",
        f"- Automated Match Score: {summary['automated_fidelity_mean']:.2f}/5 "
        f"(median {validation['automated_fidelity_median']:.1f})",
        f"- Human-Reviewed Match Score: {summary['human_fidelity_mean']:.2f}/5 "
        f"(median {validation['human_fidelity_median']:.1f})",
        f"- Mean difference (Automated Match Score - Human-Reviewed Match Score): "
        f"{validation['mean_difference']:+.2f}",
        f"- Mean absolute score difference: {validation['mean_absolute_score_difference']:.2f}",
        f"- Exact score agreement: {validation['exact_score_agreement_count']}/{summary['dataset_size']} "
        f"({validation['exact_score_agreement_percentage']:.1f}%)",
        f"- Within one point: {validation['within_one_point_count']}/{summary['dataset_size']} "
        f"({validation['within_one_point_percentage']:.1f}%)",
        f"- Rows with score difference greater than one: {validation['rows_with_score_difference_over_one']}",
        f"- Spearman (human fidelity vs G-Eval Core, reused from Step 7B): rho={spearman['rho']:.3f}, "
        f"p={spearman['p_value']:.2e}, n={spearman['n']}",
        "",
        "## Contradiction Validation",
        "",
        f"- Correctly detected: Claim-based check {contradiction['contradictions_correctly_detected']['step4']}, "
        f"NLI check {contradiction['contradictions_correctly_detected']['nli']}",
        f"- Missed by claim-based check: {contradiction['contradictions_missed_by_step4']['count']} "
        f"(rows {contradiction['contradictions_missed_by_step4']['rows']})",
        f"- Missed by NLI check: {contradiction['contradictions_missed_by_nli']['count']} "
        f"(rows {contradiction['contradictions_missed_by_nli']['rows']})",
        f"- NLI check false positives: {contradiction['nli_false_positives']['count']} "
        f"(rows {contradiction['nli_false_positives']['rows']})",
        "",
        "## Person-Level Results",
        "",
    ]
    for person in summary["person_summary"]:
        lines.append(
            f"- **{person['person_id']}** ({person['number_of_rows']} rows, consistency: "
            f"{person['consistency_label']}): Automated Match Score {person['automated_fidelity_mean']:.2f}, "
            f"Human-Reviewed Match Score {person['human_fidelity_mean']:.2f}, "
            f"contradictions {person['human_contradiction_count']}, omissions {person['human_omission_count']}"
        )
        lines.append(f"  - {person['person_final_conclusion']}")

    lines += ["", "## Representative Examples", ""]
    for example in summary["representative_examples"]:
        lines.append(
            f"- Row {example['row_id']} ({example['person_id']}): Automated Match Score "
            f"{example['automated_fidelity_score']}/5, Human-Reviewed Match Score "
            f"{example['human_fidelity_score']}/5 (Human–AI Match: "
            f"{business_match_label(example['final_fidelity_label'])}), main issue: "
            f"{business_facing_text(example['main_issue'])}. {example['explanation']}"
        )

    lines += ["", "## Final Business Findings", ""]
    lines += [f"- {finding}" for finding in summary["final_business_findings"]]

    lines += [
        "",
        "## Overall Conclusion",
        "",
        summary["overall_conclusion"],
        "",
    ]
    return "\n".join(lines)


def _style_header_row(sheet, row: int = 1) -> None:
    for cell in sheet[row]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = sheet.cell(row=row + 1, column=1).coordinate
    sheet.auto_filter.ref = f"A{row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"


def _set_column_widths(sheet, widths: dict) -> None:
    for column_letter, width in widths.items():
        sheet.column_dimensions[column_letter].width = width


def _wrap_columns(sheet, columns: list, start_row: int = 2) -> None:
    for column_letter in columns:
        for row in range(start_row, sheet.max_row + 1):
            sheet[f"{column_letter}{row}"].alignment = Alignment(wrap_text=True, vertical="top")


def _assert_no_empty_columns(sheet) -> None:
    """Guards against an accidental blank spacer column inside the used range; every
    column in a reviewer-facing sheet must carry a header and, for data sheets, values."""
    for column in range(1, sheet.max_column + 1):
        values = [sheet.cell(row=row, column=column).value for row in range(1, sheet.max_row + 1)]
        if all(value in (None, "") for value in values):
            raise FinalEvaluationError(
                f"Sheet '{sheet.title}' has a fully empty column at position {column}."
            )


def build_excel_workbook(records: list[dict], summary: dict) -> Workbook:
    workbook = Workbook()

    pairwise_sheet = workbook.active
    pairwise_sheet.title = "Pairwise Evaluation"
    pairwise_sheet.append([
        "Row ID", "Person ID", "Question", "Human Answer", "AI Answer",
        "Automated Match Score", "Human-Reviewed Match Score", "Human–AI Match", "Main Issue",
        "Final Assessment", "Claim Coverage", "Unsupported Rate", "Claim-Based Contradiction Check",
        "NLI Contradiction Check", "Human Contradiction", "Human Omission", "Human Unsupported Detail",
    ])
    for record in records:
        pairwise_sheet.append([
            record["row_id"], record["person_id"], record["question"], record["human_answer"],
            record["ai_answer"], record["automated_fidelity_score"], record["human_fidelity_score"],
            business_match_label(record["final_fidelity_label"]), business_facing_text(record["main_issue"]),
            business_facing_text(record["final_assessment_note"]),
            record["claim_coverage_rate"], record["unsupported_rate"], record["contradicted_count"] > 0,
            record["nli_contradiction_present"], record["human_contradiction"], record["human_omission"],
            record["human_unsupported_detail"],
        ])
    _style_header_row(pairwise_sheet)
    _set_column_widths(pairwise_sheet, {
        "A": 8, "B": 12, "C": 45, "D": 45, "E": 45, "F": 12, "G": 12, "H": 16, "I": 24,
        "J": 55, "K": 12, "L": 12, "M": 14, "N": 12, "O": 14, "P": 12, "Q": 16,
    })
    _wrap_columns(pairwise_sheet, ["C", "D", "E", "J"])

    person_sheet = workbook.create_sheet("Person Summary")
    person_sheet.append([
        "Person ID", "Automated Match Score (Mean)", "Human-Reviewed Match Score (Mean)", "Consistency Label",
        "Contradictions", "Omissions", "Unsupported Detail Count", "Claim Coverage",
        "Main Strength", "Main Weakness", "Final Conclusion",
    ])
    for person in summary["person_summary"]:
        coverage = person["mean_claim_coverage_rate"]
        person_sheet.append([
            person["person_id"], round(person["automated_fidelity_mean"], 2),
            round(person["human_fidelity_mean"], 2), person["consistency_label"],
            person["human_contradiction_count"], person["human_omission_count"],
            person["human_unsupported_detail_count"],
            round(coverage, 2) if coverage is not None else None,
            person["consistent_traits"], person["inconsistent_traits"], person["person_final_conclusion"],
        ])
    _style_header_row(person_sheet)
    _set_column_widths(person_sheet, {
        "A": 12, "B": 20, "C": 18, "D": 16, "E": 14, "F": 12, "G": 20, "H": 14,
        "I": 50, "J": 50, "K": 55,
    })
    _wrap_columns(person_sheet, ["I", "J", "K"])

    overall_sheet = workbook.create_sheet("Overall Summary")
    overall_sheet.append(["Metric", "Value"])
    validation = summary["validation"]
    spearman = validation["spearman_human_fidelity_vs_geval_core"]
    failure = summary["failure_evidence"]
    for label, value in [
        ("Dataset size", summary["dataset_size"]),
        ("Automated Match Score", round(summary["automated_fidelity_mean"], 2)),
        ("Human-Reviewed Match Score", round(summary["human_fidelity_mean"], 2)),
        ("Mean absolute score difference", round(validation["mean_absolute_score_difference"], 2)),
        ("Exact score agreement",
         f"{validation['exact_score_agreement_count']}/{summary['dataset_size']} "
         f"({validation['exact_score_agreement_percentage']:.1f}%)"),
        ("Within one point agreement",
         f"{validation['within_one_point_count']}/{summary['dataset_size']} "
         f"({validation['within_one_point_percentage']:.1f}%)"),
        ("Spearman rho (human fidelity vs G-Eval Core)", round(spearman["rho"], 3)),
        ("Spearman p-value", spearman["p_value"]),
        ("Spearman n", spearman["n"]),
        ("Human contradiction rate", f"{failure['human_contradiction_rate']:.1f}%"),
        ("Human omission rate", f"{failure['human_omission_rate']:.1f}%"),
        ("Human unsupported-detail rate", f"{failure['human_unsupported_detail_rate']:.1f}%"),
        ("Missing Human claim rate", f"{failure['missing_human_claim_rate']:.1f}%"),
        ("Unsupported AI claim rate", f"{failure['unsupported_ai_claim_rate']:.1f}%"),
    ]:
        overall_sheet.append([label, value])

    overall_sheet.append([None, None])
    overall_sheet.append(["Final Business Findings", None])
    for index, finding in enumerate(summary["final_business_findings"], start=1):
        overall_sheet.append([f"Finding {index}", finding])

    overall_sheet.append([None, None])
    overall_sheet.append(["Overall Conclusion", summary["overall_conclusion"]])

    _style_header_row(overall_sheet)
    _set_column_widths(overall_sheet, {"A": 32, "B": 90})
    _wrap_columns(overall_sheet, ["B"])

    for sheet in (pairwise_sheet, person_sheet, overall_sheet):
        _assert_no_empty_columns(sheet)

    return workbook
